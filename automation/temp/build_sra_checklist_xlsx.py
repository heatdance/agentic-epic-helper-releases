# -*- coding: utf-8 -*-
"""One-shot: checklist markdown-ish text -> xlsx (+ becomes -)."""
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAW = r"""# Common
# User Groups
> Permission matrix is available here https://confluence.in.devexperts.com/display/SR/Roles+Permissions+Matrix
+ User in VIP group has full access to Application
+ User in Sales Group has limited access to Application (Sales group is assigned, no additional roles added)
> Following Parts are available:  
> - Control Viewer in Read-Only mode (except Schedule A page)
> - Proposals with disabled edit, SRA script, Option Matrix, Get market Data, CSV Import, Custom Strategy
> Sales user has access to: 
> - All created Proposals, not only created by this user
> - Portfolio Analysis
> - Client Dashboard
+ User in Client Group has access only to Client Dashboard functionality
>  CD table without the Opt. Level, Sales Specialist, Sales Support fields.
>  Account/Strategy Summary doesn't have Acc. Option Level, Sales Specialist, Sales Support metrics
>  A limited version of the Mandate History panel (https://confluence.in.devexperts.com/requirements/SR/SRA-239.2)
>  CD Transactions tab without Exec Broker MPID and SEC TYPE fields.
+ User doesn't have any access to application if he is not a part of some group/doesn't have any specific role
# Advisors
+ Create new Advisor with Coverage, Market Code, market leader fields
>  IA/BA/FA
>  All fields are displayed according to user's input
+ Create new Advisor and add CEB, OYE, OYENT and some others as approved strategies
>  values should be available in tbladvisorapprovedstrategies table in DB
+ Duplicate Advisor
+ Create New Prospect Advisor
+ Convert Prospect Advisor into Client
+ Select Client Live tab
> Live tab contains live advisors only - not the entire tree structure.
+ Download the CSV file for Live tab
> The file contains only Advisors with Live status
+ Select Client Closed tab
> Closed tab contains closed advisors only - not the entire tree structure.
+ Download CSV file for Closed tab
> The file contains only Advisors with Closed status
+ Switch Client-Prospect toggle to Prospect
> Switching to Prospect disables ContractType filter
+ Change Status of Prospect Advisor to Closed
>  Closed Date field becomes editable
+ Change Status of Prospect Advisor back to Live
>  Closed Date field becomes optional and greyed out
+ Select Prospect Live tab
+ Download CSV file for Prospect Live tab
+ Select Prospect Closed tab
+ Download CSV file for Prospect Closed tab
+ Add Enterprise BillingID to Advisor
> available only if enterprise billing flag is set to YES
+ Open the assigned existing Fee schedule (Billing ID)
>  Fee page is opened in another tab, and user is navigated to the specified Schedule ID.
+ Select Market Code in Coverage Contact section
>  Market Leader and Coverage fields are autopopulated based on mapping in sradb.tblsalescoveragemapping table
+ Create new Advisors mapping
> Mapping should have unique Mapping key - Office Range combination
+ Create strategies mapping
## Advisor tab - Advisor - Contacts
+ Advisor card contains the contact section
+ New adviser's contact list is empty
+ Create a new contact and attach it to the advisor
+ Attach an existing contact to the prospect
+ The selection window displays the assigned roles for this contact
+ Change the roles of an attached contact
+ Unlink contact
## Advisor tab - Prospects - Contacts
+ Prospect card contains the contact section
+ New prospect's contact list is empty
+ Create a new contact and attach it to the prospect
+ Attach an existing contact to the prospect
+ The selection window displays the assigned roles for this contact
+ Change the roles of an attached contact
+ Unlink contact
# Account & Strategies
+ Create new account
> The data is displayed correctly in the Card and Table view
+ Market code, Coverage, Market Leader fields are read only
+ If Market Code, Coverage, and Market Leader are not present for the child Advisor, the values of the Parent Advisor should be populated
>  Each logic sources it's value on its own, meaning one field can gather information from FA, another from BA, and another from IA.
+ Open the assigned existing Fee schedule (Schedule ID)
>  Fee page is opened in another tab, and user is navigated to the specified Schedule ID.
>  Specified Fee Schedule is highlighted
>  Link from Account also highlights strategies associated with that account.
+ Add new strategy
> Add Proposal ID to the Strategy
> Only Proposal with Status Open or Won can be added.
> Add Allowed Tickers and Allowed Tickers CSC
+ Attach a file to a strategy
+ Download the attached file from the link in the Work Queue item
+ Dupllicate closed Account
# Fees
+ Create new Schedule
+ Add new Strategy
# Custodian
+ Read only
+ Download CSV
> With search criteria
> Without search
# Schedule A
+ Create Firm
+ Create Table
# Enterprise pricing table
+ Create New Enterprise pricing record
+ Update Enterprise pricing record
+ Delete Enterprise pricing record
# Contacts Tab
## General view
+ View: Email, First Name, Last Name, Phone Number, Billing Address 1, Billing Address 2, City, State, ZIP, CRD, CRM ID, Roles
+ Default sorting by Last Name (ascending)
+ Sorting available for every column except Roles
+ Sorting persists after:
>  page reload
>  switch to another tab
>  logout-login
+ "Copy text" via right-click for every column except Roles
+ Tooltips are available
+ Check the delete contact functionality with confirmation pop-up
+ Download "csv" file
>  Apply filters
>  Without filters
+ Filtering applies "AND" logic
>  accounts filter 
>  contacts filter
>  roles filter
>  global search filter
+ Filtering persists after:
>  reload page
>  switch to another tab
>  logout-login
+ General search applies to:
>  Email
>  First Name
>  Last Name
>  Phone Number
>  Billing Address 1/2 fields (highlighting is available)
+ Error message exist when:
>  retry 
>  no data
+ Delete contact
## Contact card
+ View: Back button, Contact card header, Delete button
>  Contact card header: First Name, Last Name, Email
>  Back button functionality
>  Delete button functionality with confirmation pop-up
## Contact info
+ Last name and Email - are mandatory
+ Email must be unique
+ Create a new contact
+ Update the contact
## Contact Advisors/Accounts
+ All accounts and advisors in a tree structure (Closed is not showing)
+ Trace structural relationships, but only the selected item is modified
+ Change the width of columns
+ Edit the assignments
>  single
>  several
+ Delete the assignments
+ Roles mapping (inline)
+ Selection
>  one position
>  some positions
>  all positions
>  The selection window displays the assigned roles for this contact
+ Message if there are no linked advisors/accounts yet
+ Delete contact
## Contact modal window
+ All accounts and advisors in a tree structure
+ Selection position and assigning roles
+ Searching, selection positions and assigning roles
+ Message if the output is empty
+ "Select Only" functionality
## Roles Assignments to Advisors/Accounts
+ Check that it is possible
>  to assign a single role
>  to assign several roles 
>  to edit the role/roles
>  to unassing a single role
>  to unassing several roles
>  to unassing all roles
# Work Queue
+ Create new WQ
+ Open WQ item
+ Approve WQ
+ Strategy Validation runs when Change Log contains items STRATEGY_CONFIG
>  Update an account, save but do not push to checking
>  Update a strategy, save and push item to checker
>  The WQ item contains the Strategy validation box
# Proposals
+ It is possible to create up to 10 Proposals simultaneously
+ Check CSV button (filters & without filters)
>  Proposals created in the past 365 days are downloaded. Only the latest final versions are exported. Old versions and drafts are not exported into a CSV file.
+ Create new Proposal (Regular and Custom IA) with HEC/HEP strategy, Collar focus and check Put/Call Moneyness in strategy slides and pdf
>  On HEC/HEP Collar slide (no matter the IA), in the header, after "Underlying price", there should be two new metrics:
>  Put Moneyness: XX.XX %
>  Call Moneyness: YY.XX %.
>  XX.XX and YY.YY should be strike/underlyingPrice*100. The value is rounded up to two points after the dot.
>  For XX.XX it should use strike price of Put option, for YY.YY strike price of Call option.
>  For HEC the price of underlying equity is used, for HEP the price of underlying index.
+ Create new Proposal with EFR strategy and 'UBS' strategy focus
+ Add an Option position to proposal with EFR strategy.
+ Create a Proposal with EFR Bars
>  EFR Bars slide is added to the EFR Strategy tab as an additional slide before the standard EFR P/L slide.
+ Generate a PDF with EFR Bars
+ There is an Exclude Original EFR P/L Slide toggle on the Strategy slide for EFR strategies
>  It is disabled by default
+ Exclude the PnL slide from the Strategy tab
>  The PnL slide is also excluded from the PDF
+ Create a Proposal with EFR Bars and Ticker = BRK.B
>  Options are displayed on EFR Bar graphs
+ Create CSP suite
+ Create HEC suite with UBS Advisor and download PDF
> Description slides are added before each strategy
> Max fee is calculated as Fee + 2,5%
> PDF contains only 1 Strategic Liquidation Slide
>  There is a "What to expect in different scenarios with EFR" slide
+ Include Fees flag is set to Y and disabled for Proposals with IA=UBS
+ Create Proposal with HEC suite and ML Advisor
> Fee Rate = 0.8 when IA = ML and Strategy = HEC
> Synthetic Tax Analysis is set to YES and editable
x Create Proposal with different strategies and check Disclosures pages, footers and content of the pdf (IA=ML)
>  https://confluence.in.devexperts.com/spaces/SR/pages/396602086/Proposal+PDF+Merrill+Lynch+Advisor+-+PDF+Conditional+Content
+ Create Proposal with MSSB Advisor
>  New Disclosures, Definitions, page headers
+ Create a Proposal with MSSB Advisor and EFR Strategy
>  EFR Bar slide Title, Bar Labels, and Bar grid is updated
>  EFR PnL slide new footer added
>  EFR Heatmap slide footer updated
+ Create HEP suite with multiple tickers
+ Create new Proposal with OYE strategy  with multiple tickers
+ Create new Proposal with OYENT strategy with multiple tickers
+ Create new off portal Proposal
+ Create Example Proposal
>  Types are not different from each other by functionality.
+ Create Models Proposal
+ Create a Proposal with a Heat Map Slide
> Heat Map Slide is only applicable for EFR Default strategies
+ Generate PDF with Heat Map Slide
+ Create Proposal with SNR suite
>  SNR suite contains the following strategies: Income, Uncapped Growth, Capped Growth, Principal Protection, Uncapped Principal Protection
+ Gearing Rate is displayed in the PDF metrics if Protection Type is Geared Buffer
+ If Collateral Type is not Treasury or Fixed Income ETF then Custom Dividend Yield field becomes mandatory
+ Create Proposal with 6 SNR strategies
>  For some of the strategies, fill the Custom Dividend yield manually
>  Custom Dividend yield field is moved from the General Proposal Request input panle to Strategy input panel
## Proposal Request Card
+ Fill Tax Status in General section
+ Fill Tax Status in strategy sectio
+ IF Tax Status is switched On in General section, it's overrides values of the related fields in strategies
+ Fill Tax Rate LT in General section
+ Fill Tax Rate LT in strategy section
+ IF Tax Rate LT is switched On in General section, it's overrides values of the related fields in strategies
+ Tax Rate ST in General section
+ Tax Rate ST in strategy section
+ IF Rate ST is switched On in General section, it's overrides values of the related fields in strategies
+ Fill the Allowed Tickers field in General Parameters section
+ Fill Notional value field in General Parameters section
+ Fill the Allowed Tickers field in SNR strategy section
+ Fill Notional value field in SNR strategy section
+ A note explaining this behavior should appear next to the General Parameters section title.
+ Cash Outlay field is available for SNR2-SNR5 strategies
>  Stored as percentage in sradb.proposalsversioninginput.AddCashOutlayPct
>  Range: -6.00000 – 15.00000 (for the percentage)
>  Precision: 5 decimal places for %, only integers for $
+ Input % value
>  $ value is calculated based on Notional
+ Input $ value
>  % value is calculated based on Notional
+ Create a Proposal with filled Cash outlay field.
+ Generate a PDF with Cash Outlay
>  Display only if value != 0
+ Create a Proposal with Include Fee flag unchecked
>  The Fee metric is not present on the PDF slide
>  The Fee Rate value is excluded from all the calculations
+ Open the Fee Rate tooltip
>  To open, click on Fee Rate label
+ Create a Proposal with Auto Calculated Fee Rate
>  Auto calculations enabled by default (if available)
>  Automatic Fee Rate calculation is only available if the IA is of ClientType = Client.
>  The Advisor should have assigned Billing ID with specified strategy (e.g., IA=AIP, Strategy=HEP)
+ Create a Proposal with manually set Fee Rate
+ Copy any existing Proposal
> off-Portal Proposals cannot be copied
> Fields: Notes, Date Requested, Allowed Tickers and attached PDF are not copied
>  Allowed Tickers field is in focus after copying
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with Advisor=UBS and Strategies !=EFR/HEC/SNR/CSP
> MKTGH0624U/S-3635326-X/XX
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with Advisor=UBS and the SNR Strategy
> MKTGH0824U/S-3771724-X/XX
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with Advisor=UBS and the HEC/EFR/CSP Strategy
> MKTGH0725U/S-4630918-X/XX
> where X/XX represent:
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with Advisor=ML for all Strategies
> MKTGH1125U/S-5016791 Page Number/ Total Pages.
+ Check MAP at the bottom right of the Proposal PDF for Proposals with Advisor=ML
> MAP: 8606062 (situated over the compliance stamp)
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with IA=MSSB
> MKTGH0825U/S-4738000-X/XX
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with non Custom Advisors
> MKTGH0725U/S-4650675-X/XX
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check the Compliance Stamp at the bottom right of the Proposal PDF for Proposals with Advisors= NWML
> MKTG0526-5468283-EXP0527
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check the pagination of the PDF
> Page numbers are in order
> Total number of pages value corresponds to the number of pages
### CS$ to CSC coversation
+ Covert CS$ into CSC by clicking "Validate Tickers" button
>  The valuie is taken from the field 'closePrc' of DBPROPOSALS.TBLSTOCKVALIDATION table.
## Strategic Liquidation
+ Create new non off portal MII Proposal with strategy liq flag and single Mutual Fund ticker
> One of PARWX, QLEIX, GIBIX, TAKNX, TORIX
## Work with Proposals
+ Market Code, Coverage, and Market Leader are populated when Advisor is selected
> Relationships can be found in the sradb.tlbsalestrategistlinking
+ Navigate to slide
+ Get market Data
+ Add option from option matrix
+ Call SRA script
+ Download PDF
+ Strat Liq Slices is set according to the Slices of related Strategy
+ Change the order of the Strategy slides
> Can be change by drag&drop the slide in the PROPOSAL PDF section
> Pdf should have the same slides order as on th UI
+ Attach Proposal logo
+ Delete Proposal logo
+ Upload Proposal logo
+ Switch Proposal table to displaying only Proposals with Prospect Advisors
+ Create Prospect Advisor from the Proposals with Sales role user
## Create Proposal version
>  new version can be created only for the saved version
+ Create New version from the last version
+ Create New Version from not the last version
+ Create new version from the Proposal Request tab
>  The user is redirected on the new version of the Proposal Request tab
+ The following fields cannot be changed in the new version:
>  adding a new strategy 
>  edit strategy and focus strategy 
>  IA 
+ Analyze button is disabled until at least one input field has been modified
+ When one of the input fields has been modified, the Analytics tab becomes disabled
+ Create a new version from the Analytics tab
>  The user is redirected on the new version of the Analytics tab
## PDF Generation and PDF history
+ Create several pdf's for one version by changing the content:
>  * Metrics Suite, Slices, Display Options Used, Exclude Annualized Dividend Yield
>  * Title page, Strategy, Objective, Logo
+ Clicking Download PDF on a proposal version Triggers direct download if only one file exists
+ Clicking Download PDF Opens the PDF History modal if multiple PDF files exist
+ PDF History contains PDFs from the currently selected version
>  Rows are ordered in reverse chronological order (most recent first)
+ Download the last PDF
+ Download some of the previous PDF
## Version History
+ If there is only one version in the group, the Version History icon is not shown
+ Open the version history
>  each version should contain the fields: Time & Date, Generated by, Version
+ Switch to another version via Version History button
# Client Dashboard
+ Open Account with New/Pending/Dormant/Closed status
>  Account Performance Report is opened without errors
+ Generate performance report for Account
+ Generate performance report for Strategy
+ Generate Performance report for Ticker
+ Generate Performance report with IA=ML
> Definitions page is available after Disclosures page
> Content of Definitions page according to https://confluence.in.devexperts.com/spaces/SR/pages/225030902/Dashboard+8+Performance+PDF#req-SRA-4115
> Footers on Disclosure and Performance pages https://confluence.in.devexperts.com/spaces/SR/pages/225030902/Dashboard+8+Performance+PDF#req-SRA-4116. Available on Performance page as well
> The standard oak stamp, MKTGH0825U/S-4746929-1/4.
+ Generate PDF for Account/Strategy with IA=MSSB
+ Generate PDF for Account/Strategy with Custodian = Morgan Stanley ( account starts with A.M.)
>  Disclosures are extended to two pages
>  Definitions page is added
>  Footer is added on the main report page and Strat Liq page
+ Check the Compliance Stamp for Performance Report with IA=MSSB or Custodian = Morgan Stanley at the bottom right of the Proposal PDF
> MKTGH0825U/S-4737810-X/XX
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Collateral line appears if Collateral flag is on
>  Applicable for CSP and SNR strategies
+ Check the Compliance Stamp at the bottom right of the PDF where IA != MSSB/ML
> MKTGH0825U/S-4746929-X/XX
> where X/XX represent:
> X= PAGE NUMBER
> XX= TOTAL NUMBER OF PAGES.
+ Check report generation for Benchmark
+ Check the Strategic Liquidation slide is generated for HEC an EFR strategies
> Example accounts: A.F.353 - HEC, A.F.1954 - EFR
+ All data in the PDF are aligned with the UI
> Find an Account with HEC strategy
+ Check the pagination of the PDF
> Page numbers are in order
> Total number of pages value corresponds to the number of pages
+ Check the Tracking error is displayed on UI and in PDF
> Displayed in the Summation sidebar on UI and in the PDF header
> Displayed only if strategy is MII (A.F.1265) ,HEP (A.F.2672) , OYE (A.S.3526) , OYENT (A.F.2229)
+ Switch between the Strategy Levels
> Use the dropdown under the Strategy Level header.
> A tooltip "Go to <> Strategy" appears on hover over the dropdown item
+ Switch to the Strategy Level from Account Level
> Select an Account and use the dropdown next to Account's Name
> A tooltip "Go to <> Strategy" appears on hover over the dropdown item
+ Switch to the Account Level from Strategy Level
> A tooltip "Go to Account" appears on hover over Account name
+ Hover over IA header in the Account summary
> A tooltip with full Advisor Name appears
+ Scroll Account Summary metric
> Scrolling arrows appears on hover
> Scrolling can be done by clicking the arrows or by mouse dragging
+ Expand/Shrink the Strategic Liquidation widget
+ Expand/Shrink the Trailing Performance Graph
> Square icon next to "Trailing Performance" graph header
+ Benchmark is autopopulated when strategy is HEP or MII, and Basis Risk Ticker is one of (SPY, IWM, PUT, BXM, QQQ, EFA, EEM)
## Graphs
+ Check that Trailing Performance graph is generated
+ Check that Risk vs Return graph is generated
+ Check the "70/30 Reference Asset" checkbox
>  70/30 Reference/Underlying is present in all the graphs
+ Generate the PDF with "70/30 Reference Asset" checkbox checked
+ Uncheck the "70/30 Reference Asset" checkbox
>  70/30 Reference/Underlying is excluded from all the graphs
+ Generate the PDF with "70/30 Reference Asset" checkbox unchecked
## Mandate History
>  A widget which provides historical strategy configuration for the strategy selected.
+ Check that Mandate History contain some data
>  The data is taken from sradb.strategyconfig table (not sradb.msgsra_strategyconfig)
+ Check Mandate History layout for Sales/VIP user (ROLE_CLIENT_DASHBOARD_FULL)
>  Contain the following fields: Start Date. End Date, Close Reason, Allocation %, Basis Risk Ticker, Basis Risk Ticker %, Basis Risk Upper Bound, Basis Risk Lower Bound, Customized Beta Lower Bound, Customized Beta Upper Bound, Leverage Ratio, Cash Mandate, Static Maturity, no Call Away Provision, no Cash Debit, Strategic Liquidation Flag, Billable, Notes, Allowed Tickers, Disallowed Tickers	
+ Check Mandate History layout for Client user (ROLE_CLIENT_DASHBOARD_BASIC)
>  Contain the following fields: Start Date. End Date, Allocation %, Basis Risk Ticker, Cash Mandate, Static Maturity, no Call Away Provision, no Cash Debit, Strategic Liquidation Flag, Allowed Tickers, Disallowed Tickers
## Taxlots, and Transactions
+ Check that Taxlots tab displays some data
> Select some account first (A.F.165 for example)
+ Check that Transactions tab displays some data
> Select some account first (A.F.165 for example)
## Positions
### Filters
+ Use Symbol filter
+ Filter state persists after reloading the page
### Positions table
+ Edit Positions quantity
>  Old value is displayed near the new quantity
+ Restore the original quantity
>  Restore button is displayed after clicking on the edited value again
+ The position cannot go from the negative to positive and vice versa. It cannot go past zero. Selling position will not turn into buying, buying will not turn into selling.
### Simulated Positions
+ Available with ROLE_LIVE_RISK_FULL role
+ Not available with ROLE_LIVE_RISK role
+ After editing quantity in Positions table, the simulated position is created automatically
>  The quantity of the simulated Position emulates the change in the position.
+ Automatically created position marked - "Closing #<number of the row in the symbol group>" for the positions that result in the original positions equating to zero.
+ Automatically created position marked - "Adjusting #<number of the row in the symbol goup>" for positions that result in the original positions not equating to zero.
+ Edit quantity in Simulated Position table
>  Quantity in Positions table is recalculated
+ Simulated Positions are persisted after refreshing the page
+ If the quantity in Simulated Positions is changed to "0", the simulated position is deleted.
+ Delete the Simulated Position
>  Position Quantity in Positions table is restored to original
+ Delete All simulated positions
#### Add Simulated Trade
+ Add Simulated Trade manually
+ Add multiple identical simulated trades (same side, type, expiration). They will be "joined" into one, as long as the resulting quantity does not reverse the polarity of the real position (goes past zero).
+ Add Simulated Trade that will change the side of the original Position
>  Validation is triggered with message "Quantity cannot reverse the existing position"
### Live Risk Report
+ Available with any of ROLE_LIVE_RISK role and ROLE_LIVE_RISK_FULL role
+ Analyze Positions
+ Change Experation
+ If the live risk does not match the set of selected positions anymore, this is highlighted at the bottom of the Report and Generate PDF button is disabled
+ Generate a PDF
## Accounts table
+ Use the Search bar
>  searching and highlighting by columns:
>  - accountStatus
>  - accnt
>  - strategyLevel
>  - strategy
>  - strategyFocus
>  - custodianAccnt
>  - lastName
>  - firstName
>  - institutionalAdvisor
>  - branchAdvisor
>  - financialAdvisor
+ Clear the search with xClear button
>  The button is disabled if no filter is applied and no input in General Search
+ Sort the table
>  Sorting is available for columns:
>  Strategy Focus, Custodian, Start Date, IA, BA, FA, Allocation, Mandate, Hold Start Date
+ Copy something from the Accounts table
>  Copying can be done by clicking right mouse button and using context menu - for table data
>  Success toast message should appear after copying
+ [FAILED] Long form name for IA/BA/FA is displayed when user hovers over IA/BA/FA code
+ Client full name is displayed on hover over Client Name
+ Go to the Strategy page by clicking on the line
+ Go to the Account page by clicking the link in the Account's name
+ The layout settings persists after reloading the page (inside the session)
+ The layout settings persists after logging out (between the sessions)
+ Clear Filters button clears all filters and input in the Search field
>  The button doesn't clear sorting
+ When there are new Accounts or Strategies available for display, the Notification appears on the top right side
>  To see notification: Scroll down the Accounts table after creating new Account
>  The text is:
>  New Data Available
>  Refresh the page to update the table with the latest information
+ There is a button Refresh Page on the Notification message window
+ Change the size of the columns
+ User's input in the Search bar persists after reloading the page
+ Sorting persists after reloading the page
### Lock/Freeze columns
+ Lock any column (right click on the column header)
>  Selecting “Lock Columns to the Left” locks the selected column and all columns positioned to its left.
+ Scroll the unlocked area
+ Move any column from unlock to lock area (column becomes locked)
+ Move any column from lock to unlock area (column becomes unlocked)
+ Unlock columns
## Sidebar
+ Open/Close side menu
+ Change the sequence of the columns
>  The sequence can be changed in the table or sidebar. Changes are displayed in the table and sidebar at the same time
+ Exclude some columns / all columns
>  Columns that cannot be excluded: Account, Strategy Level, Strategy Focus
+ Reset table settings to default
>  By default all columns are displayed
>  The columns are displayed in the default order
+ Check 'Auto size all columns'
>  columns are resized to such width so that they fit the content available in them.
### Filters
+ Use  Select Only filter (select one item/several items/all items)
>  Columns: Account Status, Opt. Level, Strategy Status, Strategy  Level, Tax Status, Market Code, No Call Away Provision, No Cash Debit, Strategic Liquidation
+ Use  Select - Search filter (select one item/several items/all items)
>  Columns: Strategy Focus, Basis Ticker, Custodian, IA, BA, FA, Sales Coverage, Sales Specialist, Sales Support
+ Use the Search bar in the Select - Search filter
+ IA/BA/FA filter should allow searching by code and Name
+ Use  Search Only filter (select one item/several items)
>  Columns: Account, Allowed Tickers, Client Name
+ Use  Text filter
>  Columns: Custodian Account
>  Logic is 'Starts with'
>  It is not possible to select several items in the text filter
+ Use  Numeric filter (select one item/several items)
>  Columns: Allocation %,Mandate
>  the filter has 'AND' logic
+ Use Date filter
>  Start Date, Hold Start Date columns
+ Use several filters simultaneously
+ Use filter and the Search Box simultaneously
+ Filters state persists after reloading the page
# Portfolio Analysis
## Portfolio List and Search
+ Portfolio Analysis is available for user with role ROLE_PORTFOLIO_ANALYSIS
+ Create new Portfolio
> Draft version is created in the versioning panel
+ Delete Portfolio
>  Portfolio can be deleted if the user has ROLE_ALLOWED_DELETE_RIGHTS assigned 
## Portfolio Input
+ Analyze functionality is blocked when tickers are not present
+ Add new ticker
> The field can accept characters and some symbols like . / \ - _  but not ,(comma).
+ 'Ticker Expected' error when ticker field is in focus
+ 'Quantity expected' error when CSC field is in focus
+ Bulk edit list of tickers
## Analysis table
+ Analyze Portfolio
> When at least one ticker is added to the portfolio input a button "Analyze" will be available
> The analysis table displays "Analyzing ... \r\n This may take up to 3 minutes. You will receive a notification" message in the middle of the screen until a response from the Everysk API is received
+ [FAILED] Leave the Portfolio tab during analyzing process
> Notification"<Portfolio Name> portfolio analysis is complete" will be displayed when analysis is finished
x Navigate to portfolio results from hyperlink in notification
> System will redirect to the portfolio details tab
## Reanalyze
+ Edit ticker quantity
+ Reanalyze Portfolio
+ Reanalyze increases portfolio version
+ Reanalyze is available only once per business day
> If a record is present for the current day, the user is not allowed to reanalyze the portfolio again. The reanalyze option is reset every day after 5 am CT
+ Reanalyze is disabled when draft version is present
## Copy to proposal
+ Create a new Proposal with multiple tickers
> All tickers and their quantities from the portfolio input panel are loaded into the Allowed Tickers
> New and Edit table buttons are not visible. Unlink portfolio button is present
+ New Proposal creation is not available for portfolio draft
+ New Proposal creation is not available when analysis in progress
+ Open Portfolio from linked Proposal
+ Open Proposal link from Portfolio
+ Unlink Portfolio
> All data populated in the "Allowed ticker" field will be kept and the +New / Edit table buttons will appear, but the portfolio link will be broken
+ Navigate to the analytics slide in linked Proposal
+ Generate Strategic Liquidation slide for linked Proposal
+ Portfolio link is available only for Version 1 of a given proposal ID.
# Documents
+ Documents Vault is accessible with roles ROLE_DOCUMENT , ROLE_VIP, ROLE_SALES
+ Search can be performed by:
>  File Name
>  Account Code
>  Advisor Code
>  Custodian Account
+ Use Type filter
+ Use Period filter
>  Select Only, but the available values cover the last 8 quarters period.
+ Use Account filter
+ Sort table by
>  Type
>  Period
>  Account
+ Download the document
>  Performance Report
>  Billing
"""


def parse(raw: str) -> list[dict]:
    rows: list[dict] = []
    section = ""
    subsection = ""
    subsubsection = ""
    current = None

    for line in raw.splitlines():
        s = line.rstrip()
        if not s.strip():
            continue
        if s.startswith("#### "):
            subsubsection = s[5:].strip()
            current = None
            continue
        if s.startswith("### "):
            subsubsection = s[4:].strip()
            current = None
            continue
        if s.startswith("## "):
            subsection = s[3:].strip()
            subsubsection = ""
            current = None
            continue
        if s.startswith("# "):
            section = s[2:].strip()
            subsection = ""
            subsubsection = ""
            current = None
            continue
        if s.startswith("> "):
            note = s[2:].strip()
            if current is None:
                rows.append(
                    {
                        "section": section,
                        "subsection": subsection,
                        "subsubsection": subsubsection,
                        "marker": "",
                        "item": "(section note)",
                        "details": note,
                    }
                )
                current = rows[-1]
            else:
                current["details"] = (
                    f"{current['details']}\n{note}" if current["details"] else note
                )
            continue

        if s.startswith("+ "):
            marker, text = "-", s[2:].strip()
        elif s.startswith("x "):
            marker, text = "x", s[2:].strip()
        elif s.startswith("- "):
            marker, text = "-", s[2:].strip()
        else:
            marker, text = "-", s.strip()

        current = {
            "section": section,
            "subsection": subsection,
            "subsubsection": subsubsection,
            "marker": marker,
            "item": text,
            "details": "",
        }
        rows.append(current)
    return rows


def build(rows: list[dict], out: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"

    headers = [
        "#",
        "Section",
        "Subsection",
        "Sub-subsection",
        "Status",
        "Check item",
        "Details",
        "Result",
        "Notes",
    ]
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    palette = [
        "E8F0FE",
        "EAF7EA",
        "FFF4E5",
        "F3E8FF",
        "E8F7F7",
        "FDECEC",
        "F5F5F5",
        "E8EEF7",
    ]
    section_fills: dict[str, PatternFill] = {}
    fail_fill = PatternFill("solid", fgColor="F8CBAD")

    for i, r in enumerate(rows, start=1):
        ws.append(
            [
                i,
                r["section"],
                r["subsection"],
                r["subsubsection"],
                r["marker"],
                r["item"],
                r["details"],
                "",
                "",
            ]
        )

    for col in range(1, len(headers) + 1):
        cell = ws.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        sec = row[1].value or ""
        if sec not in section_fills:
            section_fills[sec] = PatternFill(
                "solid", fgColor=palette[len(section_fills) % len(palette)]
            )
        fill = section_fills[sec]
        item = row[5].value or ""
        marker = row[4].value or ""
        if marker == "x" or "[FAILED]" in item:
            fill = fail_fill
        for cell in row:
            cell.alignment = wrap
            cell.border = thin
            cell.fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"
    widths = {
        "A": 6,
        "B": 22,
        "C": 28,
        "D": 24,
        "E": 8,
        "F": 70,
        "G": 70,
        "H": 12,
        "I": 24,
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 22
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 45

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main() -> None:
    rows = parse(RAW)
    out = Path(__file__).resolve().parent / "SRA-checklist.xlsx"
    build(rows, out)
    counts = Counter(r["marker"] for r in rows)
    print(f"saved {out}")
    print(f"rows={len(rows)}")
    print(f"markers={dict(counts)}")
    print(f"plus_remaining={sum(1 for r in rows if r['marker'] == '+')}")


if __name__ == "__main__":
    main()
